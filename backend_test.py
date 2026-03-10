import requests
import sys
import json
import io
from datetime import datetime
import openpyxl
from openpyxl import Workbook

class LogAnalysisTester:
    def __init__(self, base_url="https://log-insight-viewer.preview.emergentagent.com"):
        self.base_url = base_url
        self.api_url = f"{base_url}/api"
        self.tests_run = 0
        self.tests_passed = 0
        self.log_data_id = None

    def run_test(self, name, method, endpoint, expected_status, data=None, files=None):
        """Run a single API test"""
        url = f"{self.api_url}/{endpoint}"
        headers = {}
        
        if files is None:
            headers['Content-Type'] = 'application/json'

        self.tests_run += 1
        print(f"\n🔍 Testing {name}...")
        print(f"   URL: {url}")
        
        try:
            if method == 'GET':
                response = requests.get(url, headers=headers)
            elif method == 'POST':
                if files:
                    response = requests.post(url, files=files)
                else:
                    response = requests.post(url, json=data, headers=headers)

            success = response.status_code == expected_status
            if success:
                self.tests_passed += 1
                print(f"✅ Passed - Status: {response.status_code}")
                try:
                    response_data = response.json()
                    print(f"   Response keys: {list(response_data.keys()) if isinstance(response_data, dict) else 'Non-dict response'}")
                    return True, response_data
                except:
                    return True, {}
            else:
                print(f"❌ Failed - Expected {expected_status}, got {response.status_code}")
                try:
                    error_detail = response.json()
                    print(f"   Error: {error_detail}")
                except:
                    print(f"   Error: {response.text}")
                return False, {}

        except Exception as e:
            print(f"❌ Failed - Error: {str(e)}")
            return False, {}

    def test_root_endpoint(self):
        """Test root API endpoint"""
        success, response = self.run_test(
            "Root API Endpoint",
            "GET",
            "",
            200
        )
        return success

    def test_demo_data(self):
        """Test demo data endpoint"""
        success, response = self.run_test(
            "Demo Data Endpoint",
            "GET",
            "demo-data",
            200
        )
        
        if success and response:
            # Validate response structure
            required_fields = ['id', 'logs', 'uploaded_at', 'filename']
            if all(field in response for field in required_fields):
                print(f"   ✅ Response has all required fields")
                print(f"   📊 Demo data contains {len(response.get('logs', []))} log entries")
                self.log_data_id = response.get('id')
                
                # Validate log structure
                if response.get('logs'):
                    first_log = response['logs'][0]
                    log_fields = ['log_type', 'log_stamp', 'log_summary']
                    if all(field in first_log for field in log_fields):
                        print(f"   ✅ Log entries have correct structure")
                    else:
                        print(f"   ❌ Log entries missing fields: {[f for f in log_fields if f not in first_log]}")
                        return False
                return True
            else:
                missing_fields = [f for f in required_fields if f not in response]
                print(f"   ❌ Response missing fields: {missing_fields}")
                return False
        return success

    def create_test_excel_file(self):
        """Create a test Excel file for upload testing"""
        wb = Workbook()
        ws = wb.active
        
        # Add headers
        ws['A1'] = 'Log Type'
        ws['B1'] = 'Log Stamp'
        ws['C1'] = 'Log Summary'
        
        # Add test data
        test_logs = [
            ['ERROR', '2025-01-08 11:00:00', 'Test database connection failed'],
            ['WARNING', '2025-01-08 11:01:00', 'Test high memory usage detected'],
            ['INFO', '2025-01-08 11:02:00', 'Test user session started'],
            ['ERROR', '2025-01-08 11:03:00', 'Test API timeout occurred'],
        ]
        
        for i, log in enumerate(test_logs, 2):
            ws[f'A{i}'] = log[0]
            ws[f'B{i}'] = log[1]
            ws[f'C{i}'] = log[2]
        
        # Save to bytes
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        return excel_buffer

    def test_excel_upload(self):
        """Test Excel file upload endpoint"""
        excel_file = self.create_test_excel_file()
        
        files = {
            'file': ('test_logs.xlsx', excel_file, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        }
        
        success, response = self.run_test(
            "Excel Upload Endpoint",
            "POST",
            "upload-excel",
            200,
            files=files
        )
        
        if success and response:
            # Validate upload response
            required_fields = ['id', 'logs', 'uploaded_at', 'filename']
            if all(field in response for field in required_fields):
                print(f"   ✅ Upload response has all required fields")
                print(f"   📊 Uploaded file contains {len(response.get('logs', []))} log entries")
                if not self.log_data_id:  # Use uploaded data if demo data failed
                    self.log_data_id = response.get('id')
                return True
            else:
                missing_fields = [f for f in required_fields if f not in response]
                print(f"   ❌ Upload response missing fields: {missing_fields}")
                return False
        return success

    def test_invalid_file_upload(self):
        """Test upload with invalid file type"""
        # Create a text file instead of Excel
        text_content = "This is not an Excel file"
        files = {
            'file': ('test.txt', io.StringIO(text_content), 'text/plain')
        }
        
        success, response = self.run_test(
            "Invalid File Upload (should fail)",
            "POST",
            "upload-excel",
            400,
            files=files
        )
        return success

    def test_generate_summary(self):
        """Test AI summary generation"""
        if not self.log_data_id:
            print("❌ No log data ID available for summary generation")
            return False
            
        success, response = self.run_test(
            "Generate AI Summary",
            "POST",
            "generate-summary",
            200,
            data={"log_data_id": self.log_data_id}
        )
        
        if success and response:
            # Validate summary response
            required_fields = ['summary', 'total_logs', 'error_count', 'warning_count', 'info_count', 'log_types']
            if all(field in response for field in required_fields):
                print(f"   ✅ Summary response has all required fields")
                print(f"   📊 Summary stats - Total: {response.get('total_logs')}, Errors: {response.get('error_count')}")
                print(f"   📝 Summary length: {len(response.get('summary', ''))} characters")
                return True
            else:
                missing_fields = [f for f in required_fields if f not in response]
                print(f"   ❌ Summary response missing fields: {missing_fields}")
                return False
        return success

    def test_get_log_data(self):
        """Test retrieving specific log data"""
        if not self.log_data_id:
            print("❌ No log data ID available for retrieval test")
            return False
            
        success, response = self.run_test(
            "Get Log Data by ID",
            "GET",
            f"log-data/{self.log_data_id}",
            200
        )
        
        if success and response:
            required_fields = ['id', 'logs', 'uploaded_at', 'filename']
            if all(field in response for field in required_fields):
                print(f"   ✅ Retrieved log data has all required fields")
                return True
            else:
                missing_fields = [f for f in required_fields if f not in response]
                print(f"   ❌ Retrieved log data missing fields: {missing_fields}")
                return False
        return success

    def test_summary_with_invalid_id(self):
        """Test summary generation with invalid log data ID"""
        success, response = self.run_test(
            "Generate Summary with Invalid ID (should fail)",
            "POST",
            "generate-summary",
            404,
            data={"log_data_id": "invalid-id-12345"}
        )
        return success

def main():
    print("🚀 Starting Log Analysis Dashboard Backend Tests")
    print("=" * 60)
    
    tester = LogAnalysisTester()
    
    # Test sequence
    tests = [
        ("Root API", tester.test_root_endpoint),
        ("Demo Data", tester.test_demo_data),
        ("Excel Upload", tester.test_excel_upload),
        ("Invalid File Upload", tester.test_invalid_file_upload),
        ("Generate Summary", tester.test_generate_summary),
        ("Get Log Data", tester.test_get_log_data),
        ("Invalid Summary ID", tester.test_summary_with_invalid_id),
    ]
    
    results = {}
    for test_name, test_func in tests:
        try:
            results[test_name] = test_func()
        except Exception as e:
            print(f"❌ {test_name} failed with exception: {str(e)}")
            results[test_name] = False
            tester.tests_run += 1
    
    # Print final results
    print("\n" + "=" * 60)
    print("📊 BACKEND TEST RESULTS")
    print("=" * 60)
    
    for test_name, passed in results.items():
        status = "✅ PASSED" if passed else "❌ FAILED"
        print(f"{test_name:<30} {status}")
    
    print(f"\nOverall: {tester.tests_passed}/{tester.tests_run} tests passed")
    success_rate = (tester.tests_passed / tester.tests_run * 100) if tester.tests_run > 0 else 0
    print(f"Success Rate: {success_rate:.1f}%")
    
    return 0 if tester.tests_passed == tester.tests_run else 1

if __name__ == "__main__":
    sys.exit(main())