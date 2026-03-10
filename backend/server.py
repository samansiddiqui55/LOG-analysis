from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional
import uuid
from datetime import datetime, timezone
import io
from openpyxl import load_workbook
from emergentintegrations.llm.chat import LlmChat, UserMessage

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Define Models
class LogEntry(BaseModel):
    model_config = ConfigDict(extra="ignore")
    log_type: str
    log_stamp: str
    log_summary: str

class LogDataResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    logs: List[LogEntry]
    uploaded_at: str
    filename: str

class SummaryRequest(BaseModel):
    log_data_id: str

class SummaryResponse(BaseModel):
    summary: str
    total_logs: int
    error_count: int
    warning_count: int
    info_count: int
    log_types: dict

# Demo data
DEMO_LOGS = [
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:15:23", "log_summary": "Database connection failed: timeout after 30s"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:16:45", "log_summary": "Authentication service unreachable"},
    {"log_type": "WARNING", "log_stamp": "2025-01-08 10:17:00", "log_summary": "High memory usage detected: 85%"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:18:30", "log_summary": "Failed to process payment: invalid card"},
    {"log_type": "INFO", "log_stamp": "2025-01-08 10:19:00", "log_summary": "User session started: user_id=12345"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:20:15", "log_summary": "API rate limit exceeded for endpoint /api/users"},
    {"log_type": "WARNING", "log_stamp": "2025-01-08 10:21:30", "log_summary": "Disk space running low: 15% remaining"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:22:00", "log_summary": "SSL certificate validation failed"},
    {"log_type": "INFO", "log_stamp": "2025-01-08 10:23:00", "log_summary": "Scheduled backup completed successfully"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:24:30", "log_summary": "Queue overflow: message broker unresponsive"},
    {"log_type": "WARNING", "log_stamp": "2025-01-08 10:25:00", "log_summary": "Slow query detected: 5.2s execution time"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:26:15", "log_summary": "File upload failed: size exceeds limit"},
    {"log_type": "INFO", "log_stamp": "2025-01-08 10:27:00", "log_summary": "Cache cleared for user preferences"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:28:45", "log_summary": "Webhook delivery failed: endpoint returned 503"},
    {"log_type": "WARNING", "log_stamp": "2025-01-08 10:29:30", "log_summary": "Deprecated API version being used"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:30:00", "log_summary": "Email notification failed: SMTP error"},
    {"log_type": "INFO", "log_stamp": "2025-01-08 10:31:00", "log_summary": "New user registration: email=test@example.com"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:32:30", "log_summary": "Microservice communication timeout: order-service"},
    {"log_type": "WARNING", "log_stamp": "2025-01-08 10:33:00", "log_summary": "Session expiring in 5 minutes"},
    {"log_type": "ERROR", "log_stamp": "2025-01-08 10:34:15", "log_summary": "Data synchronization failed between nodes"},
]

@api_router.get("/")
async def root():
    return {"message": "Log Insight Viewer API"}

@api_router.post("/upload-excel", response_model=LogDataResponse)
async def upload_excel(file: UploadFile = File(...)):
    """Upload an Excel file with log data"""
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Only Excel files (.xlsx, .xls) are accepted")
    
    try:
        contents = await file.read()
        workbook = load_workbook(filename=io.BytesIO(contents))
        sheet = workbook.active
        
        logs = []
        headers = [cell.value.lower().strip() if cell.value else "" for cell in sheet[1]]
        
        # Find column indices
        type_idx = None
        stamp_idx = None
        summary_idx = None
        
        for i, header in enumerate(headers):
            if "type" in header:
                type_idx = i
            elif "stamp" in header or "time" in header or "date" in header:
                stamp_idx = i
            elif "summary" in header or "message" in header or "description" in header:
                summary_idx = i
        
        if type_idx is None or stamp_idx is None or summary_idx is None:
            raise HTTPException(
                status_code=400, 
                detail="Excel must have columns for log type, timestamp, and summary"
            )
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[type_idx] and row[stamp_idx]:
                log_stamp = str(row[stamp_idx])
                logs.append(LogEntry(
                    log_type=str(row[type_idx]).upper(),
                    log_stamp=log_stamp,
                    log_summary=str(row[summary_idx]) if row[summary_idx] else ""
                ))
        
        if not logs:
            raise HTTPException(status_code=400, detail="No valid log entries found in Excel file")
        
        # Store in MongoDB
        doc_id = str(uuid.uuid4())
        doc = {
            "id": doc_id,
            "logs": [log.model_dump() for log in logs],
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
            "filename": file.filename
        }
        await db.log_data.insert_one(doc)
        
        return LogDataResponse(
            id=doc_id,
            logs=logs,
            uploaded_at=doc["uploaded_at"],
            filename=file.filename
        )
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error processing Excel file: {e}")
        raise HTTPException(status_code=500, detail=f"Error processing file: {str(e)}")

@api_router.get("/demo-data", response_model=LogDataResponse)
async def get_demo_data():
    """Get demo log data for testing"""
    doc_id = str(uuid.uuid4())
    logs = [LogEntry(**log) for log in DEMO_LOGS]
    
    doc = {
        "id": doc_id,
        "logs": DEMO_LOGS,
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "filename": "demo_data.xlsx"
    }
    await db.log_data.insert_one(doc)
    
    return LogDataResponse(
        id=doc_id,
        logs=logs,
        uploaded_at=doc["uploaded_at"],
        filename="demo_data.xlsx"
    )

@api_router.post("/generate-summary", response_model=SummaryResponse)
async def generate_summary(request: SummaryRequest):
    """Generate AI-powered summary of log data"""
    # Fetch log data from MongoDB
    log_data = await db.log_data.find_one({"id": request.log_data_id}, {"_id": 0})
    
    if not log_data:
        raise HTTPException(status_code=404, detail="Log data not found")
    
    logs = log_data["logs"]
    
    # Calculate statistics
    total_logs = len(logs)
    error_count = sum(1 for log in logs if log["log_type"] == "ERROR")
    warning_count = sum(1 for log in logs if log["log_type"] == "WARNING")
    info_count = sum(1 for log in logs if log["log_type"] == "INFO")
    
    # Count by type
    log_types = {}
    for log in logs:
        log_type = log["log_type"]
        log_types[log_type] = log_types.get(log_type, 0) + 1
    
    # Prepare error summaries for AI
    error_summaries = [log["log_summary"] for log in logs if log["log_type"] == "ERROR"]
    warning_summaries = [log["log_summary"] for log in logs if log["log_type"] == "WARNING"]
    
    # Generate AI summary
    try:
        api_key = os.environ.get("EMERGENT_LLM_KEY")
        if not api_key:
            # Fallback to basic summary if no API key
            summary = f"Analysis of {total_logs} log entries:\n\n"
            summary += f"• {error_count} errors detected\n"
            summary += f"• {warning_count} warnings detected\n"
            summary += f"• {info_count} info entries\n\n"
            if error_summaries:
                summary += "Top error types:\n"
                for err in error_summaries[:5]:
                    summary += f"  - {err}\n"
        else:
            chat = LlmChat(
                api_key=api_key,
                session_id=f"log-analysis-{request.log_data_id}",
                system_message="You are a DevOps expert analyzing system logs. Provide concise, actionable insights."
            ).with_model("openai", "gpt-5.2")
            
            prompt = f"""Analyze these system logs and provide a brief, actionable summary:

Total Logs: {total_logs}
Errors: {error_count}
Warnings: {warning_count}
Info: {info_count}

Error Messages:
{chr(10).join(error_summaries[:10])}

Warning Messages:
{chr(10).join(warning_summaries[:5])}

Provide:
1. A 2-3 sentence overview of the system health
2. Top 3 critical issues that need immediate attention
3. Recommended actions to resolve the issues

Keep the response under 250 words."""
            
            user_message = UserMessage(text=prompt)
            summary = await chat.send_message(user_message)
            
    except Exception as e:
        logger.error(f"AI summary generation failed: {e}")
        # Fallback summary
        summary = f"**Log Analysis Report**\n\n"
        summary += f"Total entries analyzed: {total_logs}\n"
        summary += f"Critical errors: {error_count}\n"
        summary += f"Warnings: {warning_count}\n"
        summary += f"Information logs: {info_count}\n\n"
        if error_summaries:
            summary += "**Key Issues Identified:**\n"
            for i, err in enumerate(error_summaries[:5], 1):
                summary += f"{i}. {err}\n"
    
    return SummaryResponse(
        summary=summary,
        total_logs=total_logs,
        error_count=error_count,
        warning_count=warning_count,
        info_count=info_count,
        log_types=log_types
    )

@api_router.get("/log-data/{log_id}", response_model=LogDataResponse)
async def get_log_data(log_id: str):
    """Get specific log data by ID"""
    log_data = await db.log_data.find_one({"id": log_id}, {"_id": 0})
    
    if not log_data:
        raise HTTPException(status_code=404, detail="Log data not found")
    
    return LogDataResponse(**log_data)

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
